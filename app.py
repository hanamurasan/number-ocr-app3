import io,re,json,base64,urllib.request,urllib.error
from pathlib import Path
import cv2,joblib,numpy as np,pandas as pd,streamlit as st
from PIL import Image
from skimage.feature import hog
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_from_string,column_index_from_string,get_column_letter

B=Path(__file__).parent
MODELS=joblib.load(B/'digit_models.joblib')
SLOTS=[(0,39),(35,74),(82,124)]
REPO='hanamurasan/number-ocr-app3'
CORRECTIONS_PATH='corrections.json'
st.set_page_config(page_title='学習する測定値OCR',page_icon='🔢',layout='wide')
st.title('🔢 修正結果から学習する測定値OCR')
st.caption('加工前写真専用。表示はXX.X、範囲は20.0～99.9です。間違いを直して「修正結果を学習」を押すと次回判定へ反映します。')

def timekey(n):
 m=re.search(r'_(\d+(?:\.\d+)?)s(?:\.[^.]+)?$',n,re.I)
 return (0,float(m.group(1))) if m else (1,n.lower())

def align(rgb):
 h,w=rgb.shape[:2];cx=round(w*.505);cy=round(h*.544);size=max(20,round(min(h,w)*.072));u=size/30
 b=(max(0,int(cx-48*u)),max(0,int(cy-33*u)),min(w,int(cx-1*u)),min(h,int(cy-3*u)))
 r=rgb[b[1]:b[3],b[0]:b[2]]
 if r.size==0:return None,b
 return cv2.resize(cv2.cvtColor(r,cv2.COLOR_RGB2GRAY),(126,72)),b

def feature(q):
 q=cv2.resize(q,(32,48));q=cv2.createCLAHE(2,(4,4)).apply(q)
 x=hog(q,orientations=9,pixels_per_cell=(8,8),cells_per_block=(2,2),block_norm='L2-Hys').astype(np.float32)
 return x/(np.linalg.norm(x)+1e-8)

def github_headers(token=''):
 h={'Accept':'application/vnd.github+json','X-GitHub-Api-Version':'2022-11-28','User-Agent':'number-ocr-app'}
 if token:h['Authorization']=f'Bearer {token}'
 return h

def github_token():
 try:return str(st.secrets.get('GITHUB_TOKEN',''))
 except:return ''

def load_remote_corrections():
 url=f'https://api.github.com/repos/{REPO}/contents/{CORRECTIONS_PATH}'
 try:
  req=urllib.request.Request(url,headers=github_headers(github_token()))
  with urllib.request.urlopen(req,timeout=10) as res:data=json.loads(res.read().decode('utf-8'))
  content=base64.b64decode(data['content']).decode('utf-8')
  return json.loads(content),data.get('sha')
 except Exception:
  try:return json.loads((B/CORRECTIONS_PATH).read_text(encoding='utf-8')),None
  except:return [],None

def save_remote_corrections(items,sha):
 token=github_token()
 if not token:return False,'StreamlitのSecretsにGITHUB_TOKENが設定されていません。'
 url=f'https://api.github.com/repos/{REPO}/contents/{CORRECTIONS_PATH}'
 body={'message':'Add OCR correction','content':base64.b64encode(json.dumps(items,ensure_ascii=False,separators=(',',':')).encode()).decode(),'branch':'main'}
 if sha:body['sha']=sha
 req=urllib.request.Request(url,data=json.dumps(body).encode(),headers={**github_headers(token),'Content-Type':'application/json'},method='PUT')
 try:
  with urllib.request.urlopen(req,timeout=15) as res:res.read()
  return True,'修正結果を保存しました。次回判定から反映します。'
 except urllib.error.HTTPError as e:
  return False,f'GitHub保存エラー {e.code}: {e.read().decode("utf-8",errors="ignore")[:200]}'
 except Exception as e:return False,f'保存エラー: {e}'

CORRECTIONS,REMOTE_SHA=load_remote_corrections()
if 'session_corrections' not in st.session_state:st.session_state.session_corrections=[]

def corrected_probability(pos,f,base_prob,classes):
 samples=[x for x in CORRECTIONS+st.session_state.session_corrections if int(x.get('position',-1))==pos]
 if not samples:return base_prob
 sims=[]
 for x in samples:
  z=np.asarray(x['feature'],dtype=np.float32);z=z/(np.linalg.norm(z)+1e-8)
  sims.append((float(np.dot(f,z)),str(x['label'])))
 sims.sort(reverse=True)
 best_sim,best_label=sims[0]
 # Only a genuinely close prior correction may override; otherwise retain model.
 if best_sim>=0.92:
  out=base_prob*0.45
  idx=np.where(classes==best_label)[0]
  if len(idx):out[idx[0]]+=0.55
  return out/out.sum()
 return base_prob

def recognize(rgb):
 a,b=align(rgb)
 if a is None:return '',0,b,[],None,[]
 out='';details=[];cs=[];features=[]
 for pos,(x1,x2) in enumerate(SLOTS):
  f=feature(a[2:70,x1:x2]);features.append(f);m=MODELS[pos];pr=m.predict_proba([f])[0];pr=corrected_probability(pos,f,pr,m.classes_);order=np.argsort(pr)[::-1];d=str(m.classes_[order[0]]);out+=d;cs.append(float(pr[order[0]]));details.append([(str(m.classes_[j]),float(pr[j])) for j in order[:4]])
 value=out[:2]+'.'+out[2]
 try:valid=20.0<=float(value)<=99.9
 except:valid=False
 return value,(float(np.mean(cs)) if valid else 0.0),b,details,a,features

def valid_value(x):
 if not re.fullmatch(r'\d{2}\.\d',x.strip()):return False
 return 20.0<=float(x)<=99.9

def cellok(x):return bool(re.fullmatch(r'[A-Za-z]{1,3}[1-9][0-9]*',x.strip()))

st.write(f'保存済み補正データ：{len(CORRECTIONS)}桁')
if not github_token():st.info('この版を永続学習させるにはStreamlitのSecretsへGITHUB_TOKENを設定します。未設定でも、この起動中は修正が反映されます。')
excel=st.file_uploader('1. 入力先Excel',type=['xlsx'])
fs=sorted(st.file_uploader('2. 加工前写真を選択',type=['png','jpg','jpeg','webp'],accept_multiple_files=True) or [],key=lambda f:timekey(f.name));rows=[]
if fs:st.info('時間順：'+' → '.join(f.name for f in fs))
for i,f in enumerate(fs):
 try:rgb=np.array(Image.open(io.BytesIO(f.getvalue())).convert('RGB'))
 except Exception as e:st.error(f'{f.name}を開けません：{e}');continue
 predicted,conf,b,details,a,features=recognize(rgb);mark=rgb.copy();cv2.rectangle(mark,(b[0],b[1]),(b[2],b[3]),(0,255,0),2)
 left,right=st.columns([1,2]);left.image(mark,caption=f.name,width='stretch')
 value=right.text_input('認識結果（間違っていれば修正）',predicted,key=f'value_{i}_{f.name}')
 right.write(f'平均確率 {conf:.2f}')
 if conf<.55:right.warning('確率が低いため確認してください。')
 if right.button('修正結果を学習',key=f'learn_{i}_{f.name}'):
  if not valid_value(value):right.error('20.0～99.9のXX.X形式で入力してください。')
  elif value==predicted:right.info('認識結果と同じなので追加学習は不要です。')
  else:
   digits=value.replace('.','');new=[]
   for pos,(label,feat) in enumerate(zip(digits,features)):
    new.append({'position':pos,'label':label,'feature':np.round(feat,6).tolist(),'source':f.name})
   st.session_state.session_corrections.extend(new)
   merged=CORRECTIONS+st.session_state.session_corrections
   ok,msg=save_remote_corrections(merged,REMOTE_SHA)
   (right.success if ok else right.warning)(msg)
 with right.expander('候補を見る'):
  if a is not None:right.image(a,caption='認識に使った数値部分',width='stretch')
  for k,z in enumerate(details):right.write(f'{k+1}桁目：'+', '.join(f'{x}({p:.2f})' for x,p in z))
 rows.append({'filename':f.name,'value':value,'confidence':round(conf,3)})

if st.session_state.session_corrections:
 payload=json.dumps(CORRECTIONS+st.session_state.session_corrections,ensure_ascii=False).encode('utf-8')
 st.download_button('補正データのバックアップをダウンロード',payload,'corrections.json','application/json')
if rows:
 st.dataframe(pd.DataFrame(rows),width='stretch')
 if excel:
  wb=load_workbook(io.BytesIO(excel.getvalue()));sn=st.selectbox('入力シート',wb.sheetnames);start=st.text_input('開始セル','C4')
  if cellok(start):
   letters,r0=coordinate_from_string(start.upper());col=column_index_from_string(letters);ws=wb[sn];bad=[]
   for j,x in enumerate(rows):
    try:cell=ws.cell(r0+j,col,float(x['value']));cell.number_format='0.0'
    except:bad.append(x['filename'])
   if bad:st.error('数値にできない画像：'+'、'.join(bad))
   else:
    out=io.BytesIO();wb.save(out);end=f'{get_column_letter(col)}{r0+len(rows)-1}';st.success(f'{start.upper()}:{end}へ入力しました');st.download_button('入力済みExcelをダウンロード',out.getvalue(),f'{Path(excel.name).stem}_入力済み.xlsx')
