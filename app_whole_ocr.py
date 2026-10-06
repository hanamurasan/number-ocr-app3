import io,re
from pathlib import Path
import cv2,numpy as np,pandas as pd,streamlit as st
from PIL import Image
from rapidocr import RapidOCR
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_from_string,column_index_from_string,get_column_letter

st.set_page_config(page_title='測定値・文字列OCR',page_icon='🔢',layout='wide')
st.title('🔢 「XX.X」を文字列のまま読み取ってExcelへ')
st.caption('桁別分類モデルは使いません。画像内の数値文字列全体をOCRし、20.0～99.9の候補を採用します。')

@st.cache_resource
def get_ocr():
    return RapidOCR()
OCR=get_ocr()

def timekey(name):
    m=re.search(r'_(\d+(?:\.\d+)?)s(?:\.[^.]+)?$',name,re.I)
    return (0,float(m.group(1))) if m else (1,name.lower())

def clean_text(text):
    table=str.maketrans({'O':'0','o':'0','I':'1','l':'1','|':'1','S':'5','s':'5','B':'8','，':'.','。':'.',',':'.'})
    return str(text).translate(table)

def candidates_from_text(text,score,source):
    t=clean_text(text)
    out=[]
    # Prefer an explicit decimal expression.
    for m in re.finditer(r'(?<!\d)(\d{2})\s*\.\s*(\d)(?!\d)',t):
        v=float(m.group(1)+'.'+m.group(2))
        if 20.0<=v<=99.9:out.append((v,float(score)+0.20,source,text))
    # Also accept a valid expression followed by OCR noise/digits, e.g. 49.70.
    for m in re.finditer(r'(\d{2})\s*\.\s*(\d)',t):
        v=float(m.group(1)+'.'+m.group(2))
        if 20.0<=v<=99.9:out.append((v,float(score)+0.12,source,text))
    # If the decimal point disappeared, insert it before the third digit.
    digits=''.join(re.findall(r'\d',t))
    if len(digits)>=3:
        for i in range(len(digits)-2):
            v=float(digits[i:i+2]+'.'+digits[i+2])
            if 20.0<=v<=99.9:
                bonus=0.05 if i==0 else 0.0
                out.append((v,float(score)+bonus,source,text))
    return out

def read_variant(name,img,weight):
    # Upscaling is only preprocessing; OCR still reads the complete string.
    up=cv2.resize(img,None,fx=3,fy=3,interpolation=cv2.INTER_CUBIC)
    result=OCR(up)
    rows=[]
    if result is not None and result.txts is not None:
        for text,score in zip(result.txts,result.scores):
            rows.append((str(text),float(score)))
    cands=[]
    for text,score in rows:cands.extend(candidates_from_text(text,min(1.0,score*weight),name))
    return rows,cands

def recognize(rgb):
    bgr=cv2.cvtColor(rgb,cv2.COLOR_RGB2BGR);h,w=bgr.shape[:2]
    # Full image plus a broad center search area. No fixed character slots.
    variants=[('画像全体',bgr,1.00)]
    center=bgr[int(h*.38):int(h*.68),int(w*.20):int(w*.78)]
    variants.append(('中央領域',center,1.08))
    # Contrast-enhanced center as a third independent reading.
    lab=cv2.cvtColor(center,cv2.COLOR_BGR2LAB);l,a,b=cv2.split(lab);l=cv2.createCLAHE(2.0,(8,8)).apply(l)
    enhanced=cv2.cvtColor(cv2.merge([l,a,b]),cv2.COLOR_LAB2BGR)
    variants.append(('中央・コントラスト補正',enhanced,1.04))
    raw=[];allc=[]
    for name,img,weight in variants:
        rows,cands=read_variant(name,img,weight);raw.append((name,rows));allc.extend(cands)
    if not allc:return '',0.0,raw,center,[]
    # Combine repeated votes for the same value and retain the strongest evidence.
    grouped={}
    for v,s,src,text in allc:
        grouped.setdefault(v,[]).append((s,src,text))
    ranked=[]
    for v,evidence in grouped.items():
        scores=sorted([x[0] for x in evidence],reverse=True)
        combined=scores[0]+0.12*max(0,len(evidence)-1)+0.05*sum(scores[1:3])
        ranked.append((combined,v,evidence))
    ranked.sort(reverse=True)
    best=ranked[0]
    confidence=min(1.0,best[0])
    return f'{best[1]:.1f}',confidence,raw,center,ranked[:5]

def cellok(x):return bool(re.fullmatch(r'[A-Za-z]{1,3}[1-9][0-9]*',x.strip()))

excel=st.file_uploader('1. 入力先Excel',type=['xlsx'])
files=sorted(st.file_uploader('2. 加工前写真を選択',type=['png','jpg','jpeg','webp'],accept_multiple_files=True) or [],key=lambda f:timekey(f.name))
if files:st.info('時間順：'+' → '.join(f.name for f in files))
rows=[]
for i,f in enumerate(files):
    try:rgb=np.array(Image.open(io.BytesIO(f.getvalue())).convert('RGB'))
    except Exception as e:st.error(f'{f.name}を開けません：{e}');continue
    value,conf,raw,center,ranked=recognize(rgb)
    left,right=st.columns([1,2]);left.image(rgb,caption=f.name,width='stretch')
    value=right.text_input('認識結果',value,key=f'{i}_{f.name}')
    right.write(f'OCR候補の確度：{conf:.2f}')
    if not value:right.error('20.0～99.9の数値候補を取得できませんでした。')
    elif conf<.70:right.warning('候補の一致が弱いため確認してください。')
    with right.expander('OCRの読み取り内容'):
        right.image(cv2.cvtColor(center,cv2.COLOR_BGR2RGB),caption='中央の探索領域',width='stretch')
        for name,items in raw:
            right.write(name+'：'+('、'.join(f'{t} ({s:.2f})' for t,s in items) if items else '文字なし'))
        if ranked:
            right.write('数値候補：'+', '.join(f'{v:.1f}' for _,v,_ in ranked))
    rows.append({'filename':f.name,'value':value,'confidence':round(conf,3)})

if rows:
    st.dataframe(pd.DataFrame(rows),width='stretch')
    if excel:
        wb=load_workbook(io.BytesIO(excel.getvalue()));sheet=st.selectbox('入力シート',wb.sheetnames);start=st.text_input('開始セル','C4')
        if cellok(start):
            letters,r0=coordinate_from_string(start.upper());col=column_index_from_string(letters);ws=wb[sheet];bad=[]
            for j,row in enumerate(rows):
                try:
                    v=float(row['value'])
                    if not 20.0<=v<=99.9:raise ValueError
                    cell=ws.cell(r0+j,col,v);cell.number_format='0.0'
                except:bad.append(row['filename'])
            if bad:st.error('有効な数値にできない画像：'+'、'.join(bad))
            else:
                out=io.BytesIO();wb.save(out);end=f'{get_column_letter(col)}{r0+len(rows)-1}'
                st.success(f'{start.upper()}:{end}へ入力しました')
                st.download_button('入力済みExcelをダウンロード',out.getvalue(),f'{Path(excel.name).stem}_入力済み.xlsx')
